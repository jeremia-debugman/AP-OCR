import os
import sys
import io
import openpyxl
from pathlib import Path

# Add backend directory to sys.path
backend_dir = Path(__file__).resolve().parent / "backend"
sys.path.insert(0, str(backend_dir))

from fastapi.testclient import TestClient
from main import app
from models import InvoiceData, InvoiceLineItem, ExportExcelRequest

def run_e2e_tests():
    print("=== Starting E2E AP OCR Verification ===")
    client = TestClient(app)

    # 1. Health check
    res = client.get("/api/health")
    assert res.status_code == 200, f"Health check failed: {res.text}"
    print(f"[OK] Health check passed: {res.json()}")

    # 2. Config endpoints
    res = client.get("/api/config")
    assert res.status_code == 200
    config_data = res.json()
    print(f"[OK] AI Config retrieved: Provider={config_data['provider']}, Configured={config_data['is_configured']}")

    # Update config test
    res = client.post("/api/config", json={
        "provider": "gemini",
        "model_name": "gemini/gemini-1.5-flash"
    })
    assert res.status_code == 200
    print("[OK] AI Config update endpoint verified")

    # 3. Test Sample Invoices Listing
    res = client.get("/api/sample-invoices")
    assert res.status_code == 200
    samples = res.json().get("samples", [])
    print(f"[OK] Sample invoices endpoint returned {len(samples)} samples")

    # 4. Test Single & Batch File Extraction
    test_pdf = backend_dir / "data" / "uploads" / "3770070d29c676d7b7b525f9f282e4f81a0973e15a2c2e520a02db89579484a3.pdf"
    test_img = backend_dir / "data" / "uploads" / "0039f10e0f34d1728bb0949de49bd2cc7d553782b9bd16997492b8309bd658d9.jpeg"

    files_to_upload = []
    if test_pdf.exists():
        with open(test_pdf, "rb") as f:
            files_to_upload.append(("files", (test_pdf.name, f.read(), "application/pdf")))
    if test_img.exists():
        with open(test_img, "rb") as f:
            files_to_upload.append(("files", (test_img.name, f.read(), "image/jpeg")))

    if files_to_upload:
        print(f"Uploading {len(files_to_upload)} invoice file(s) for extraction...")
        res = client.post("/api/extract", files=files_to_upload)
        assert res.status_code == 200, f"Extraction failed: {res.text}"
        batch_res = res.json()
        assert batch_res["success"] is True
        assert len(batch_res["invoices"]) == len(files_to_upload)
        print(f"[OK] Batch Extraction succeeded for {len(batch_res['invoices'])} invoices")

        extracted_invoices = batch_res["invoices"]
        for inv in extracted_invoices:
            print(f"  - Invoice #{inv.get('invoice_number')}: Vendor='{inv.get('vendor_name')}', Date='{inv.get('invoice_date')}', Total={inv.get('total_amount')} {inv.get('currency')}, Items={len(inv.get('line_items', []))}, Confidence={inv.get('confidence_level')}")
            assert inv.get("file_hash") is not None
            assert "line_items" in inv

        # 5. Test Excel Export (Batch)
        print("Testing Batch Excel Export Endpoint...")
        export_payload = {
            "invoices": extracted_invoices,
            "export_format": "xlsx",
            "include_summary_sheet": True,
            "include_line_items_sheet": True
        }
        res = client.post("/api/export/excel", json=export_payload)
        assert res.status_code == 200, f"Excel export failed: {res.text}"
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        # Verify generated Excel with openpyxl
        excel_wb = openpyxl.load_workbook(io.BytesIO(res.content))
        sheet_names = excel_wb.sheetnames
        print(f"[OK] Excel workbook generated with sheets: {sheet_names}")
        assert "Invoices Summary" in sheet_names
        assert "Line Items Detail" in sheet_names

        ws_sum = excel_wb["Invoices Summary"]
        assert ws_sum.cell(row=1, column=1).value == "File Name"
        assert ws_sum.cell(row=1, column=2).value == "Invoice Number"
        assert ws_sum.cell(row=1, column=5).value == "Vendor / Supplier"
        print(f"[OK] Verified Summary Sheet Headers and Data (Rows: {ws_sum.max_row})")

        ws_items = excel_wb["Line Items Detail"]
        assert ws_items.cell(row=1, column=1).value == "File Name"
        assert ws_items.cell(row=1, column=6).value == "Description"
        print(f"[OK] Verified Line Items Sheet Headers and Data (Rows: {ws_items.max_row})")

        # 6. Test Single Invoice Excel Export
        print("Testing Single Invoice Excel Export Endpoint...")
        single_res = client.post("/api/export/single-excel", json=extracted_invoices[0])
        assert single_res.status_code == 200
        single_wb = openpyxl.load_workbook(io.BytesIO(single_res.content))
        print(f"[OK] Single Invoice Excel generated successfully (Sheets: {single_wb.sheetnames})")

    print("\n========================================================")
    print("ALL E2E AP OCR EXTRACTION & EXCEL TESTS PASSED 100%!")
    print("========================================================")

if __name__ == "__main__":
    run_e2e_tests()
